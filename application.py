import sys
import os
import threading
import time
import tkinter as tk
from tkinter import filedialog, scrolledtext, messagebox, ttk, Listbox, END
from datetime import datetime
from pathlib import Path
import logging
import winsound

from config import Config
from separating_robots import split_report_by_robot
from analyzer import analyze_report, APIKeyExpiredError
from support_bot_analyzer import analyze_support_report, save_support_report_with_summary
from pivot_generator import (
    generate_pivot_report as _generate_pivot_report_backend,
    create_full_report as _create_full_report_backend,
)

import warnings

warnings.filterwarnings('ignore')


# ===== ВЫСОКОУРОВНЕВЫЕ ФУНКЦИИ (используют pivot_generator и analyzer, без дублирования) =====
def run_pivot_only(input_excel_path: str, progress_callback=None):
    """
    Создаёт отчёт только со сводными таблицами (режим 'Только сводные').

    Возвращает: (путь_к_выходному_файлу или None, путь_к_лог_файлу или None)
    """
    if progress_callback:
        progress_callback("📈 Генерация сводных таблиц...")
    return _generate_pivot_report_backend(input_excel_path)


def run_errors_only(input_excel_path: str, agent_name: str, progress_callback=None):
    """
    Запускает только AI‑анализ ошибок.
    """
    if progress_callback:
        progress_callback("🔍 AI‑анализ ошибок...")
    errors_df, log_file = analyze_report(input_excel_path, agent_name)

    if errors_df is None or errors_df.empty:
        return None, log_file

    if progress_callback:
        progress_callback("💾 Сохранение файла с ошибками...")
    timestamp = datetime.now().strftime('%d.%m.%Y_%H-%M')
    original_name = Path(input_excel_path).stem
    agent_filename = agent_name.replace(" ", "_").replace("-", "_").lower()

    Path("processed").mkdir(exist_ok=True)
    output_path = f"processed/errors_{agent_filename}_{timestamp}_{original_name}.xlsx"

    # Сохраняем через pandas (заголовки колонок в первой строке)
    errors_df.to_excel(output_path, index=False)

    # Открываем для форматирования
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = load_workbook(output_path)
    ws = wb.active
    ws.title = "Отчёт по ошибкам"

    # Заголовки колонок уже на строке 1, просто форматируем их
    headers = list(errors_df.columns)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Автоширина колонок
    for col_idx in range(1, len(headers) + 1):
        max_len = 0
        for row in range(1, ws.max_row + 1):
            cell_value = ws.cell(row=row, column=col_idx).value
            if cell_value:
                max_len = max(max_len, len(str(cell_value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

    # Границы для всех ячеек с данными
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))
    for row in range(1, ws.max_row + 1):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border

    wb.save(output_path)

    return output_path, log_file


def run_support_categorization(input_excel_path: str, progress_callback=None):
    """
    Запускает AI-категоризацию входящих обращений.
    Добавляет в конец Excel колонки 'Категория' и 'Тип_ошибки'.
    """
    if progress_callback:
        progress_callback("🔍 AI-категоризация входящих обращений...")
    support_df, log_file = analyze_support_report(input_excel_path)

    if support_df is None or support_df.empty:
        return None, log_file

    if progress_callback:
        progress_callback("💾 Сохранение файла с категориями...")
    timestamp = datetime.now().strftime('%d.%m.%Y_%H-%M')
    original_name = Path(input_excel_path).stem

    Path("processed").mkdir(exist_ok=True)
    output_path = f"processed/support_{timestamp}_{original_name}.xlsx"

    save_support_report_with_summary(support_df, output_path)

    return output_path, log_file


def run_full_report(input_excel_path: str, agent_name: str, logger: logging.Logger | None = None,
                    progress_callback=None):
    """
    Создаёт полный отчёт (AI‑анализ + сводные таблицы в одном файле).
    """
    if logger is None:
        logger = logging.getLogger(f"full_report_{datetime.now().strftime('%d.%m.%Y_%H-%M')}")
        logger.setLevel(logging.INFO)

    if progress_callback:
        progress_callback("🔍 AI‑анализ ошибок...")
    errors_df, log_file = analyze_report(input_excel_path, agent_name)

    if errors_df is None or errors_df.empty:
        return None, log_file

    if progress_callback:
        progress_callback("📊 Добавление сводных таблиц...")
    full_report_path = _create_full_report_backend(
        original_path=input_excel_path,
        errors_df=errors_df,
        agent_name=agent_name,
        logger=logger,
    )

    return full_report_path, log_file


# ===== КЛАСС ДЛЯ ПЕРЕНАПРАВЛЕНИЯ STDOUT =====
class TextRedirector:
    def __init__(self, text_widget):
        self.text_widget = text_widget

    def write(self, string):
        self.text_widget.after_idle(self._write, string)

    def _write(self, string):
        self.text_widget.insert(tk.END, string)
        self.text_widget.see(tk.END)

    def flush(self):
        pass


# ===== GUI-приложение для обработки Excel-отчётов =====
class ExcelProcessorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Обработчик Excel-отчётов (AI анализ + сводные)")
        self.root.geometry("750x750")
        self.root.resizable(True, True)

        self.file_path = tk.StringVar()
        self.report_type = tk.StringVar(value="pivot")
        self.selected_agent = tk.StringVar(value="")
        self.cancel_event = threading.Event()
        self.processing_thread = None
        self.last_output_path = None
        self.last_log_path = None
        self.start_time = None
        self.timer_running = False

        # Переменные для очереди
        self.queue_files = []
        self.queue_running = False
        self.is_batch_mode = False

        # ===== Фрейм для кнопки подготовки отчётов =====
        frame_prepare = tk.Frame(root)
        frame_prepare.pack(pady=15, padx=10, fill=tk.X)

        btn_split = tk.Button(
            frame_prepare,
            text="🤖 Подготовить отчёты к обработке (разделить по роботам)",
            command=self.split_files
        )
        btn_split.pack(side=tk.LEFT)

        # ===== Основная строка выбора файла =====
        frame_top = tk.Frame(root)
        frame_top.pack(pady=5, padx=10, fill=tk.X)

        lbl_file = tk.Label(frame_top, text="Файл отчёта (report.xlsx):")
        lbl_file.pack(side=tk.LEFT)

        entry_file = tk.Entry(frame_top, textvariable=self.file_path, width=50)
        entry_file.pack(side=tk.LEFT, padx=5)

        btn_browse = tk.Button(frame_top, text="Обзор...", command=self.browse_file)
        btn_browse.pack(side=tk.LEFT)

        # ===== Остальные блоки =====

        # Блок выбора типа отчёта
        frame_type = tk.LabelFrame(root, text="Тип отчёта")
        frame_type.pack(pady=5, padx=10, fill=tk.X)

        tk.Radiobutton(
            frame_type,
            text="📈 Только сводные таблицы",
            variable=self.report_type,
            value="pivot",
            anchor="w",
            justify="left",
        ).pack(side=tk.TOP, anchor="w")

        tk.Radiobutton(
            frame_type,
            text="📊 Только ошибки (AI анализ)",
            variable=self.report_type,
            value="errors",
            anchor="w",
            justify="left",
        ).pack(side=tk.TOP, anchor="w")

        tk.Radiobutton(
            frame_type,
            text="🚀 Полный отчёт (AI + сводные)",
            variable=self.report_type,
            value="full",
            anchor="w",
            justify="left",
        ).pack(side=tk.TOP, anchor="w")

        tk.Radiobutton(
            frame_type,
            text="📞 Категоризация входящих обращений (AI)",
            variable=self.report_type,
            value="support",
            anchor="w",
            justify="left",
        ).pack(side=tk.TOP, anchor="w")

        # Блок выбора агента для AI‑анализа
        frame_agent = tk.LabelFrame(root, text="Агент для AI‑анализа")
        frame_agent.pack(pady=5, padx=10, fill=tk.X)

        agents = list(Config.AGENTS.keys())
        if agents:
            self.selected_agent.set(agents[0])

        tk.Label(frame_agent, text="Агент:").pack(side=tk.LEFT)
        self.agent_menu = tk.OptionMenu(frame_agent, self.selected_agent, *agents)
        self.agent_menu.pack(side=tk.LEFT, padx=5)

        # Кнопки управления
        btn_frame = tk.Frame(root)
        btn_frame.pack(pady=5)

        self.btn_run = tk.Button(
            btn_frame,
            text="Запустить обработку",
            command=self.run_processing,
            bg="#4CAF50",
            fg="white",
            padx=10,
            pady=5,
        )
        self.btn_run.pack(side=tk.LEFT, padx=5)

        self.btn_cancel = tk.Button(btn_frame, text="Отмена", command=self.cancel_processing,
                                    bg="#f44336", fg="white", padx=10, pady=5, state=tk.DISABLED)
        self.btn_cancel.pack(side=tk.LEFT, padx=5)

        self.btn_open = tk.Button(
            btn_frame,
            text="Открыть отчёт",
            command=self.open_last_file,
            state=tk.DISABLED,
            padx=10,
            pady=5,
        )
        self.btn_open.pack(side=tk.LEFT, padx=5)

        # Блок очереди файлов
        frame_queue = tk.LabelFrame(root, text="Очередь файлов")
        frame_queue.pack(pady=5, padx=10, fill=tk.X, side=tk.TOP)

        btn_add_queue = tk.Button(frame_queue, text="➕ Добавить файлы в очередь", command=self.add_files_to_queue)
        btn_add_queue.pack(side=tk.LEFT, padx=5, pady=5)

        btn_clear_queue = tk.Button(frame_queue, text="🗑 Очистить очередь", command=self.clear_queue)
        btn_clear_queue.pack(side=tk.LEFT, padx=5, pady=5)

        self.queue_listbox = Listbox(frame_queue, height=4, selectmode=tk.MULTIPLE)
        self.queue_listbox.pack(fill=tk.X, padx=5, pady=5)

        self.btn_start_queue = tk.Button(
            frame_queue,
            text="▶ Запустить очередь",
            command=self.start_queue,
            bg="#2196F3",
            fg="white",
            padx=10,
            pady=5,
            state=tk.DISABLED
        )
        self.btn_start_queue.pack(side=tk.LEFT, padx=5, pady=5)

        # Панель прогресса
        progress_frame = tk.Frame(root)
        progress_frame.pack(pady=5, padx=10, fill=tk.X)

        self.stage_label = tk.Label(progress_frame, text="", anchor=tk.W, fg="blue")
        self.stage_label.pack(fill=tk.X)

        self.progress_bar = ttk.Progressbar(progress_frame, mode='indeterminate', length=400)
        self.progress_bar.pack(pady=5, fill=tk.X)

        self.timer_label = tk.Label(progress_frame, text="⏱️ Время: 00:00", anchor=tk.W)
        self.timer_label.pack(fill=tk.X)

        # Консоль вывода
        self.console = scrolledtext.ScrolledText(root, wrap=tk.WORD, height=15)
        self.console.pack(padx=10, pady=5, fill=tk.BOTH, expand=True)

        # Статус
        self.status_label = tk.Label(root, text="Готов к работе", anchor=tk.W)
        self.status_label.pack(fill=tk.X, padx=10, pady=5)

        self.original_stdout = sys.stdout

    def browse_file(self):
        filename = filedialog.askopenfilename(
            title="Выберите Excel-файл",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if filename:
            if not filename.lower().endswith('.xlsx'):
                messagebox.showerror("Ошибка", "Выбранный файл не является файлом Excel (.xlsx)")
                return
            self.file_path.set(filename)
            self.console.delete(1.0, tk.END)
            self.status_label.config(text=f"Выбран файл: {filename}")

    def add_files_to_queue(self):
        """Добавляет выбранные файлы в очередь"""
        filenames = filedialog.askopenfilenames(
            title="Выберите Excel-файлы для очереди",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        for f in filenames:
            if f not in self.queue_files:
                self.queue_files.append(f)
                self.queue_listbox.insert(END, f)
        self.btn_start_queue.config(state=tk.NORMAL if self.queue_files else tk.DISABLED)

    def clear_queue(self):
        """Очищает очередь"""
        self.queue_files.clear()
        self.queue_listbox.delete(0, END)
        self.btn_start_queue.config(state=tk.DISABLED)

    def start_queue(self):
        """Запускает обработку очереди"""
        if not self.queue_files:
            messagebox.showwarning("Очередь пуста", "Добавьте файлы в очередь.")
            return
        if self.queue_running:
            messagebox.showinfo("Обработка уже запущена", "Дождитесь завершения текущей очереди.")
            return

        self.queue_running = True
        self.btn_start_queue.config(state=tk.DISABLED)
        self.btn_run.config(state=tk.DISABLED)

        self.processing_thread = threading.Thread(target=self._process_queue_thread, daemon=True)
        self.processing_thread.start()

    def _process_queue_thread(self):
        """Поток для последовательной обработки файлов очереди"""
        total = len(self.queue_files)
        self.root.after(0, lambda: self.console.insert(END, f"📋 Запущена очередь из {total} файлов\n"))
        self.root.after(0, self.console.see, END)

        for idx, file_path in enumerate(self.queue_files, 1):
            if self.cancel_event.is_set():
                self.root.after(0, lambda: self.console.insert(END, "⏹ Очередь прервана пользователем.\n"))
                break

            self.root.after(0, lambda: self.console.insert(END, f"\n{'=' * 50}\n"))
            self.root.after(0, lambda: self.console.insert(END,
                                                           f"📄 Обработка файла {idx}/{total}: {Path(file_path).name}\n"))
            self.root.after(0, self.console.see, END)

            # Устанавливаем текущий файл
            self.file_path.set(file_path)

            # Запускаем обработку (синхронно)
            self.is_batch_mode = True
            self._process_single_file()
            self.is_batch_mode = False
            self.root.after(0, lambda: self.console.insert(END, f"✅ Файл {idx}/{total} обработан!\n"))
            self.root.after(0, self.console.see, END)

        self.queue_running = False
        self.root.after(0, lambda: self.btn_start_queue.config(state=tk.NORMAL if self.queue_files else tk.DISABLED))
        self.root.after(0, lambda: self.btn_run.config(state=tk.NORMAL))

        # Финальное уведомление
        winsound.PlaySound("SystemDefault", winsound.SND_ALIAS)
        self.root.after(0, lambda: self.console.insert(END, f"\n🎉 Очередь из {total} файлов полностью завершена!\n"))
        self.root.after(0, self.console.see, END)
        self.root.after(0, lambda: messagebox.showinfo("Очередь завершена", f"Обработано {total} файлов."))

    def run_processing(self):
        file_path = self.file_path.get().strip()
        if not file_path:
            messagebox.showwarning("Предупреждение", "Сначала выберите файл.")
            return

        self.processing_thread = threading.Thread(target=self._process_single_file, daemon=True)
        self.processing_thread.start()

    def _process_single_file(self):
        """Обрабатывает текущий файл из self.file_path (синхронно)"""
        file_path = self.file_path.get().strip()
        if not file_path:
            return

        self.console.delete(1.0, tk.END)
        sys.stdout = TextRedirector(self.console)

        self.cancel_event.clear()
        self.btn_run.config(state=tk.DISABLED)
        self.btn_cancel.config(state=tk.NORMAL)
        self.btn_open.config(state=tk.DISABLED)
        self.status_label.config(text="Обработка выполняется...")
        self.stage_label.config(text="🚀 Подготовка...")
        self.start_time = time.time()
        self.timer_running = True
        self._update_timer()
        self.progress_bar.start(10)

        self.last_output_path = None
        self.last_log_path = None

        report_type = self.report_type.get()
        agent_name = self.selected_agent.get().strip()

        print(f"➡️ Выбран тип отчёта: {report_type}")
        if report_type in ("errors", "full") and not agent_name:
            print("❌ Не выбран агент для AI‑анализа.")
            return

        def stage_callback(msg):
            self.root.after(0, self._update_stage, msg)

        try:
            if report_type == "pivot":
                output_path, log_path = run_pivot_only(file_path, progress_callback=stage_callback)
            elif report_type == "errors":
                output_path, log_path = run_errors_only(file_path, agent_name, progress_callback=stage_callback)
            elif report_type == "support":
                output_path, log_path = run_support_categorization(file_path, progress_callback=stage_callback)
            else:
                output_path, log_path = run_full_report(file_path, agent_name, progress_callback=stage_callback)

            if output_path:
                print(f"✅ Готово! Отчёт сохранён в файле:\n{output_path}")
                if not self.is_batch_mode:  # звук только для одиночной обработки
                    winsound.PlaySound("SystemAsterisk", winsound.SND_ALIAS)
                self.last_output_path = str(Path(output_path).absolute())
            else:
                print("❌ Не удалось создать отчёт.")

            if log_path:
                print(f"📋 Лог‑файл:\n{log_path}")
                self.last_log_path = log_path

        except APIKeyExpiredError:
            print(f"⛔ Ошибка: истёк или недействителен API‑ключ.")
        except Exception as e:
            print(f"❌ Ошибка: {e}")

        finally:
            self.root.after(100, self._processing_finished)

    def cancel_processing(self):
        if self.processing_thread and self.processing_thread.is_alive():
            self.cancel_event.set()
            self.status_label.config(text="Отмена запрошена...")
            self.btn_cancel.config(state=tk.DISABLED)

    def _update_timer(self):
        if self.timer_running:
            elapsed = int(time.time() - self.start_time)
            mins, secs = divmod(elapsed, 60)
            self.timer_label.config(text=f"⏱️ Время: {mins:02d}:{secs:02d}")
            self.root.after(1000, self._update_timer)

    def _update_stage(self, stage_text):
        self.stage_label.config(text=stage_text)

    def open_last_file(self):
        path = self.last_output_path
        if not path:
            messagebox.showinfo("Открыть отчёт", "Сначала сформируйте отчёт.")
            return
        if not Path(path).exists():
            messagebox.showerror("Открыть отчёт", f"Файл не найден:\n{path}")
            return
        try:
            os.startfile(str(path))
        except Exception as e:
            messagebox.showerror("Открыть отчёт", f"Не удалось открыть файл:\n{e}")

    def _processing_finished(self):
        sys.stdout = self.original_stdout
        self.btn_run.config(state=tk.NORMAL)
        self.btn_cancel.config(state=tk.DISABLED)
        self.progress_bar.stop()
        self.timer_running = False

        if self.last_output_path:
            self.btn_open.config(state=tk.NORMAL)
        else:
            self.btn_open.config(state=tk.DISABLED)

        if self.cancel_event.is_set():
            self.status_label.config(text="Обработка отменена.")
            self.stage_label.config(text="❌ Отменено")
        else:
            self.status_label.config(text="Обработка завершена.")
            elapsed = int(time.time() - self.start_time)
            mins, secs = divmod(elapsed, 60)
            self.timer_label.config(text=f"⏱️ Время: {mins:02d}:{secs:02d} (завершено)")
            self.stage_label.config(text="✅ Готово")

    def split_files(self):
        """Запускает разделение отчёта по роботам."""
        filename = filedialog.askopenfilename(
            title="Выберите Excel-файл для разделения",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
        )
        if not filename:
            return

        self.console.delete(1.0, tk.END)
        sys.stdout = TextRedirector(self.console)

        try:
            self.status_label.config(text="Разделение файла...")
            self.stage_label.config(text="⏳ Разделение по роботам...")
            self.progress_bar.start(10)

            threading.Thread(target=self._split_thread, args=(filename,), daemon=True).start()

        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось разделить файл:\n{e}")
            sys.stdout = self.original_stdout
            self.progress_bar.stop()
            self.status_label.config(text="Готов к работе")
            self.stage_label.config(text="")

    def _split_thread(self, filename):
        """Поток для разделения файла."""
        try:
            created_files = split_report_by_robot(filename, output_dir="split")
            if created_files:
                print(f"✅ Разделение завершено. Создано файлов: {len(created_files)}")
                for f in created_files:
                    print(f"   - {f}")

                winsound.PlaySound("SystemAsterisk", winsound.SND_ALIAS)
                messagebox.showinfo("Разделение завершено", f"Создано {len(created_files)} файла в папке 'split'")
            else:
                print("⚠️ Не найдено строк для указанных роботов.")
                messagebox.showwarning("Разделение", "Не найдено строк для указанных роботов.")
        except Exception as e:
            print(f"❌ Ошибка при разделении: {e}")
            messagebox.showerror("Ошибка", f"Ошибка разделения:\n{e}")
        finally:
            self.root.after(0, self._split_finished)

    def _split_finished(self):
        """Завершает разделение – сбрасывает статус."""
        sys.stdout = self.original_stdout
        self.progress_bar.stop()
        self.status_label.config(text="Готов к работе")
        self.stage_label.config(text="")


if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelProcessorApp(root)
    root.mainloop()
